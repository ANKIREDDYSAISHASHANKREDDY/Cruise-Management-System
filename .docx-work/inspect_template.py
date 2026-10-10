import zipfile, os
from docx import Document
f='PBL Final Documentation templet.docx'
z=zipfile.ZipFile(f)
for n in z.namelist():
    if n.startswith('word/media/'):
        z.extract(n, '.docx-work/images')
d=Document(f)
print('inline_shapes',len(d.inline_shapes),'paragraphs',len(d.paragraphs),'sections',len(d.sections))
for i,p in enumerate(d.paragraphs):
    drawings=len(p._p.xpath('.//w:drawing'))
    breaks=len(p._p.xpath('.//w:br'))
    if drawings or breaks:
        print('STRUCT',i,'drawings',drawings,'breaks',breaks,'text',repr(p.text))
for i,s in enumerate(d.sections):
    print('PAGE',i+1,'orientation',s.orientation,'size',s.page_width,s.page_height,'margins',s.left_margin,s.right_margin,s.top_margin,s.bottom_margin)
