from copy import deepcopy
from pathlib import Path
import zipfile, hashlib
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from PIL import Image, ImageDraw, ImageFont

ROOT = Path.cwd()
TMP = ROOT / '.docx-work'
SRC = ROOT / 'PBL Final Documentation templet.docx'
WORK = TMP / 'working-with-content.docx'
DRAFT = TMP / 'Harborline-PBL-draft.docx'
ART = TMP / 'generated-figures'
ART.mkdir(parents=True, exist_ok=True)
doc = Document(str(SRC))

def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise ValueError('Missing paragraph: ' + prefix)

def replace(p, text, style=None, preserve=True):
    first = next((r for r in p.runs if r.text or r._r.rPr is not None), None)
    rpr = deepcopy(first._r.rPr) if preserve and first is not None and first._r.rPr is not None else None
    for child in list(p._p):
        if child.tag != qn('w:pPr'):
            p._p.remove(child)
    if style:
        p.style = doc.styles[style]
    run = p.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)
    return p

def add_after(anchor, text, style='Body Text', exemplar=None, size=None):
    newp = deepcopy(exemplar._p) if exemplar is not None else OxmlElement('w:p')
    for child in list(newp):
        if child.tag != qn('w:pPr'):
            newp.remove(child)
    anchor._p.addnext(newp)
    p = Paragraph(newp, anchor._parent)
    if style:
        p.style = doc.styles[style]
    run = p.add_run(text)
    if size is not None:
        run.font.size = Pt(size)
    return p

def add_block(anchor, texts, style='Body Text', exemplar=None, size=None, move_section=False):
    sect = None
    if move_section and anchor._p.pPr is not None and anchor._p.pPr.sectPr is not None:
        sect = anchor._p.pPr.sectPr
        anchor._p.pPr.remove(sect)
    cursor = anchor
    made = []
    for text in texts:
        cursor = add_after(cursor, text, style, exemplar, size)
        made.append(cursor)
    if sect is not None:
        made[-1]._p.get_or_add_pPr().append(sect)
    return made

def remove(p):
    p._p.getparent().remove(p._p)

def is_picture(p, rid):
    return any(x.get(qn('r:embed')) == rid for x in p._p.xpath('.//a:blip'))

# Locate source slots before insertion.
cover_title = find('Title (20 size)')
cover_fields = {
    'in The Department of XXXX': 'in the Department of [Department Name]',
    'Course name with Course code': '[Course Name] ([Course Code])',
    'Roll.no: NAME': 'Roll No.: [Roll Number] — [Student Name]',
    'XXXX(faculty name)': '[Faculty Guide Name]',
    'Department of Electronics and Communication Engineering': '[Department Name]',
    'FEB - 2025.': '[Submission Month Year]'
}
paper_title = find('Title (Font size 24)')
paper_note = find('Note: Maintain Document format')
abstract_h = find('Abstract')
figures_h = find('List of Figures')
tables_h = find('List of Tables')
toc_h = find('Table of Contents')
intro_h = find('Introduction')
intro = find('[Note: introduction')
method_h = find('METHODOLOGY')
method_sub = find('Selecting a Template')
method_body = find('[For example] The three subtypes')
fig1 = next(p for p in doc.paragraphs if is_picture(p, 'rId10'))
experiments_h = find('EXPERIMENTS')
placeholders = [p for p in doc.paragraphs if p.text.strip() == '[For Examples]']
experiments_sub, results_sub, conclusion_sub, references_sub = placeholders
experiments_body = find('[For example] The primary goal')
fig2 = next(p for p in doc.paragraphs if is_picture(p, 'rId11'))
evaluation_h = find('Keep your text and graphic files separate')
results_h = find('RESULTS')
results_body = find('[For example] Further')
fig3 = next(p for p in doc.paragraphs if is_picture(p, 'rId12'))
tables_image = next(p for p in doc.paragraphs if is_picture(p, 'rId13'))
conclusion_h = find('CONCLUSION and FUTURE WORK')
conclusion_body = find('The In this paper')
references_h = find('References')
reference_paras = [p for p in doc.paragraphs if p.style.name == 'references']
body_ex = intro

# Cover title is project-specific; identity fields await the user's reply.
replace(cover_title, 'Harborline Cruise Management Booking and Smart Operations System')
replace(paper_title, 'Harborline Cruise Management Booking and Smart Operations System', 'paper title', False)
for p in list(doc.paragraphs):
    value = p.text.strip()
    if value in cover_fields:
        replace(p, cover_fields[value], preserve=True)
remove(paper_note)

# Abstract page, figure/table lists, and contents. Move section breaks to the last new line on each page.
abstract = (
    'Harborline is a role-based cruise management and booking prototype that connects passenger reservations with onboard and ship-operation workflows. '
    'Passengers can browse sample sailings, reserve a cabin, assign one dining seat to each ticketed guest across capacity-limited meal batches, and review itinerary notices. '
    'Crew can view shifts, record attendance, accept assigned faults or guest-service work, and log progress. Captains and managers can set up sailings, manage scoped staff permissions, control inventory, update port calls, maintain equipment, and lead muster workflows. '
    'A captain-reviewed disruption-recovery process links a delay or cancellation to affected booking parties, passenger choices, and claimable crew follow-ups. '
    'The application uses an HTML, CSS, and JavaScript interface, a same-origin Node.js HTTP API, and a local SQLite database. Server-side roles, sailing scopes, transactional booking, and activity logs support accountable workflows. '
    'Stripe Checkout can be enabled with server credentials; without them, bookings run in demo mode and take no payment. The project verification record dated 6 October 2026 reports 13 passing automated tests, but no external passenger or crew acceptance study, live vessel connection, or production deployment is recorded. '
    'This report explains the system, its development method, scenario-based evaluation, results, and implementation boundaries.'
)
abstract_ps = add_block(abstract_h, [abstract, 'Keywords: cruise management; passenger booking; ship operations; Adaptive Software Engineering.'],
                        exemplar=body_ex, size=10.5, move_section=True)
for p in abstract_ps:
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)

fig_ps = add_block(figures_h, [
    'Figure 1. Harborline architecture and configured integration boundary' + chr(9) + '7',
    'Figure 2. Role workspaces and key passenger and crew workflows' + chr(9) + '8',
    'Figure 3. Captain-led dispatch and recovery control flow' + chr(9) + '9'
], exemplar=body_ex, size=10.5, move_section=True)
for p in fig_ps:
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.45), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

table_ps = add_block(tables_h, [
    'Table 1. Implemented workflows by role' + chr(9) + '9',
    'Table 2. Prototype capabilities and integration boundaries' + chr(9) + '9'
], exemplar=body_ex, size=10.5, move_section=True)
for p in table_ps:
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.45), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

toc_ps = add_block(toc_h, [
    '1  Introduction' + chr(9) + '6',
    '2  Methodology' + chr(9) + '6',
    '   a. Development method and adaptive cycle' + chr(9) + '6',
    '   b. System architecture and data model' + chr(9) + '7',
    '   c. Authentication and operational accountability' + chr(9) + '7',
    '3  Experiments' + chr(9) + '8',
    '   a. Role-based operational scenarios' + chr(9) + '8',
    '   b. Verification method and evidence limits' + chr(9) + '8',
    '4  Results' + chr(9) + '9',
    '   a. Delivered workflows and recorded outcomes' + chr(9) + '9',
    '5  Conclusion and future work' + chr(9) + '10',
    'References' + chr(9) + '10'
], exemplar=body_ex, size=10.5, move_section=True)
for p in toc_ps:
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.25 if p.text.startswith('   ') else 0.03)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.45), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

# Introduction: problem, objectives, roles, project-specific evolution, and honest scope.
replace(intro_h, 'Introduction')
intro_parts = [
    'Cruise operations span a continuous journey from choosing a sailing to returning ashore. A reservation must represent the voyage, cabin, and the number of guests travelling together. Before embarkation, passengers need clear booking status and a usable boarding process. Once onboard, dining capacity, guest requests, port-call changes, maintenance findings, and safety drills involve different departments. If each activity is recorded in a separate or static tool, staff may repeat information or act on an out-of-date status, while passengers may see a change without a clear way to respond.',
    'Harborline was developed as a local academic prototype to connect those guest and operator workflows in one role-aware system. The project evolved through iterative requirements: server-backed sign-in and saved bookings; dining batches with a seat assignment for every ticketed guest; controls over who can change inventory; and captain-created operational work that crew members can accept and log. A later operations requirement was to make a fault report actionable: it needs a sailing, location, priority, responsible department, owner, and recorded progress rather than a static example card.',
    'The goal is to demonstrate an end-to-end management flow, not to reproduce a cruise line system of record. The passenger workspace supports sample-sailing discovery, cabin selection, reservations, dining allocation, itinerary notices, onboard plans, and service cases. The crew workspace exposes assigned shifts, attendance, work history, fault and service tasks, and selected maintenance actions. Captains and managers can configure sailings and port calls, apply sailing-scoped staff permissions, oversee equipment and muster activity, review disruption impacts, and inspect logged changes. Each view is backed by a local API and SQLite database rather than browser-only sample state [1], [2].',
    'The system has four principal user roles. Passengers manage their own booking and guest allocations. Crew members see tasks within their permitted sailings and record acceptance, progress, and completion. Captains initiate or approve operational actions, including fault dispatch, itinerary updates, muster events, and disruption recovery. Managers configure inventory and sailing access, review operations, and handle administrative decisions. The interface and API enforce these roles separately; a user does not gain authority merely by navigating to a different page.',
    'A particular design contribution is the captain-reviewed disruption workflow. For a delayed or cancelled port call, the application prepares an impact review for confirmed booking parties and linked shore plans. An authorized captain or manager reviews and approves the plan before passengers receive in-app notices. Passenger responses are recorded, and crew can claim follow-up tasks. The workflow links operational change, guest impact, and staff action in one audit trail. It is a distinctive capability in this prototype; no claim is made that it is new to the cruise industry.',
    'The implementation boundary is important. Sample sailings, fares, cabins, and some operational records are seeded. The default booking flow is a demonstration and takes no money unless the server is configured for Stripe Checkout. Camera scanning depends on browser Code 39 support, with manual and keyboard-scanner alternatives. In-app notices do not send email or SMS. The local offline queue covers selected actions, but it is not guaranteed emergency communications. There is no live connection to cruise inventory, identity, crew scheduling, port-agent, maintenance-sensor, or safety systems. Harborline therefore serves as a reviewable software increment for further stakeholder validation rather than an operational vessel service [1], [4].'
]
replace(intro, intro_parts[0], 'Body Text', False)
cursor = intro
for text in intro_parts[1:]:
    cursor = add_after(cursor, text, exemplar=body_ex)

# Methodology: ASE loop, architecture, data, access control, and audit.
replace(method_h, 'Methodology')
replace(method_sub, 'Development Method and Adaptive Cycle', 'Heading 2', False)
method_parts = [
    'The project follows an Adaptive Software Engineering (ASE) loop supported by short Agile-style increments. Speculation established a working vision and recorded uncertainty about passenger inventory, onboard capacity, crew authority, safety requirements, external data ownership, and payment. Collaboration turned user-directed changes into role-specific workflows and acceptance criteria. Learning came from implementing those flows against a local server and reviewing the resulting behavior and recorded automated checks. Adaptation then changed the backlog and prototype, including guest-level dining assignments, sailing-scoped staff rights, captain-issued faults, crew work logs, and port-disruption recovery [3].',
    'This is a development record, not evidence from a completed user study. The sample backlog contains illustrative scenarios, and no passenger or crew participant quotes or formal acceptance sessions are recorded. Product decisions in this report are therefore described as implemented requirements and local verification, not as statistically validated usability findings.'
]
replace(method_body, method_parts[0], 'Body Text', False)
cursor = method_body
for text in method_parts[1:]:
    cursor = add_after(cursor, text, exemplar=body_ex)

arch_h = add_after(cursor, 'System Architecture and Data Model', 'Heading 2', method_sub)
arch_texts = [
    'The prototype uses a same-origin browser client built with HTML, CSS, and JavaScript, served by a small Node.js HTTP server. The server exposes authenticated JSON routes and stores users, sessions, sailings, cabins, bookings, dining batches and seats, shift and task records, permission scopes, port calls, safety events, maintenance schedules, service cases, and action logs in SQLite. This keeps the academic build simple to start while providing server-side persistence and transaction boundaries. Cabin reservations use collision protection and idempotency keys; dining assignments must match the booking party size, use available seats on the same sailing, and fit the capacity of each selected batch [1], [2].',
    'Figure 1 shows the implemented local architecture alongside optional integrations. Stripe is an adapter that operates only when server-side keys and a signed webhook secret are configured. The development default uses demo mode. The client can queue selected actions locally while disconnected, then submit them for server verification after reconnection. The queue is limited and device-local; it does not replace a durable fleet-wide synchronization service.'
]
cursor = add_after(arch_h, arch_texts[0], exemplar=body_ex)
for text in arch_texts[1:]:
    cursor = add_after(cursor, text, exemplar=body_ex)

auth_h = add_after(cursor, 'Authentication and Operational Accountability', 'Heading 2', method_sub)
auth_text = (
    'Passwords are stored as scrypt hashes, and authenticated sessions use random tokens stored by hash with HttpOnly and SameSite cookies. '
    'API handlers check role, account ownership, and sailing-level permission before exposing or changing records. Staff permissions distinguish actions such as managing a sailing, editing dining inventory, dispatching faults, viewing a crew job board, scanning boarding, and managing muster. '
    'Consequential operations append actor-attributed history, including inventory adjustments, fault acceptance and work, access changes, boarding scans, maintenance completion, port updates, and disruption decisions. '
    'These controls are implemented in the local prototype but have not received an independent security review or an operator audit-retention review [1], [2].'
)
cursor = add_after(auth_h, auth_text, exemplar=body_ex)

# Experiments are scenario-based; the report distinguishes the stored test record from user research.
replace(experiments_h, 'Experiments')
replace(experiments_sub, 'Role-Based Operational Scenarios', 'Heading 2', False)
experiment_parts = [
    'Evaluation uses repeatable scenarios that exercise the principal roles and server-backed state. In the passenger scenario, a signed-in user selects a sample sailing, reserves an available cabin, and checks that a second request cannot claim an already reserved cabin. The passenger then assigns one dining seat to each guest in the booking. Each guest can be assigned to a different batch; blocked, filled, and available seats are shown, and the server rejects missing guest assignments, duplicate seats, blocked seats, or a batch over capacity.',
    'In the crew and captain scenario, a captain or manager creates a fault with sailing, deck or location, department, priority, and task details. A crew member with access to that sailing can accept the work, start it, and record resolution. The action history identifies the reporter and crew activity. Staff-access settings constrain which sailings and actions each account can manage. A separate boarding scenario checks ticket references against the authorized sailing manifest, rejects repeat scans, and records the operator who checked in the booking.',
    'For a port disruption, the captain or manager selects a delayed or cancelled call, reviews affected confirmed booking parties and linked shore plans, then approves the recovery. Only after approval are in-app notices created and passenger response options recorded. Claimable follow-up tasks are opened for crew and their progress is logged. Safety and preventive-maintenance scenarios verify that authorized staff can open a drill or emergency muster, account for ticketed guests at a configured station, and complete scheduled equipment service with a note and next due date. Figure 2 summarizes the user responsibilities covered by these scenarios.'
]
replace(experiments_body, experiment_parts[0], 'Body Text', False)
cursor = experiments_body
for text in experiment_parts[1:]:
    cursor = add_after(cursor, text, exemplar=body_ex)

replace(evaluation_h, 'Verification Method and Evidence Limits', 'Heading 2', False)
evaluation_texts = [
    'The repository test plan defines passenger, crew, captain, and manager acceptance checks and records an existing verification run dated 6 October 2026: 13 automated tests passed, comprising 7 domain tests and 6 HTTP integration tests. The documented integration cases cover authentication and authorization, reservation collisions and idempotency, SQLite-backed workflows, feedback review, crew operations, and registration validation. Syntax checks and a browser walkthrough are also recorded in that project document [4].',
    'These results are evidence for the prototype code and its seeded local workflows only. No performance or load benchmark, live payment transaction, camera-device compatibility matrix, external passenger or crew study, vessel-system integration, or production security assessment is recorded. This report does not treat scripted scenarios as stakeholder acceptance.'
]
cursor = add_after(evaluation_h, evaluation_texts[0], exemplar=body_ex)
for text in evaluation_texts[1:]:
    cursor = add_after(cursor, text, exemplar=body_ex)

# Results and boundaries.
replace(results_h, 'Results')
replace(results_sub, 'Delivered Workflows and Recorded Outcomes', 'Heading 2', False)
result_parts = [
    'The local build demonstrates a connected passenger and operations flow. A passenger account can register and sign in, browse seeded cruise options, select a cabin, and keep a reservation in the server database. Transactional checks protect cabin inventory from conflicting reservation requests. The dining workflow models seat status by batch and keeps the number of guest assignments equal to the booking party size. Onboard requests, feedback, itinerary notices, and service cases are also attached to signed-in accounts.',
    'Crew and ship operations are more than a static job list. Captains and managers can create and dispatch faults for sailings they control; crew can claim eligible jobs, update progress, and close work with a resolution note. Shift attendance, swaps, preventive maintenance, port-call status, staff permissions, boarding scans, and muster accountability have server routes and history records. The manager and captain views bring these operational records together. The disruption-recovery flow adds a human approval step before notices and follow-up tasks are created. Figure 3 illustrates this dispatch-to-closure control path.'
]
replace(results_body, result_parts[0], 'Body Text', False)
cursor = results_body
for text in result_parts[1:]:
    cursor = add_after(cursor, text, exemplar=body_ex)
add_after(tables_image,
    'Tables 1 and 2 distinguish working local workflows from integration boundaries. The strongest result is traceability across roles: a configured permission controls who may act, the server persists the decision, and subsequent crew or passenger responses remain linked to the originating sailing or booking. The prototype still depends on seeded data and local infrastructure, so this result demonstrates coherent workflow behavior rather than fleet-wide operating effectiveness.',
    exemplar=body_ex)

# Conclusion and project-source references.
replace(conclusion_h, 'Conclusion and Future Work')
remove(conclusion_sub)
conclusion_parts = [
    'Harborline demonstrates a role-based cruise-management prototype in which booking, dining allocation, boarding, guest service, crew work, maintenance, port status, and safety accountability share a persistent local backend. Its operational model gives captains and managers control over sailing-scoped actions while allowing crew to accept and log assigned jobs. Dining assignments are made per ticketed guest across capacity-limited batches. The captain-reviewed recovery workflow ties a port change to affected bookings, passenger choices, and crew follow-up, creating a concrete path from decision to resolution.',
    'The work remains an academic local MVP. The next iteration should validate priorities with representative passengers, crew, captains, and managers; connect authoritative sailing and identity sources; complete production payment, refund, notification, and manifest flows; define schedule qualifications and labor rules; test offline recovery and device support; and establish backups, audit retention, accessibility review, deployment controls, and independent security assessment. Muster station mappings and emergency actions must be configured and validated by vessel operators before any operational use. Progress should be measured through task completion and error rates in consented usability sessions, alongside reliability and security checks, rather than through assumptions about user acceptance.'
]
replace(conclusion_body, conclusion_parts[0], 'Body Text', False)
cursor = conclusion_body
for text in conclusion_parts[1:]:
    cursor = add_after(cursor, text, exemplar=body_ex)
replace(references_h, 'References')
remove(references_sub)
refs = [
    'Harborline project repository. README: implemented features, local setup, payment and scanner configuration, and production boundaries.',
    'Harborline project repository. server.js and app.js: Node.js API, SQLite persistence, authorization, booking, crew, and operations workflows.',
    'Harborline project repository. Architecture, Data Model and Risk Register; ASE and Agile Project Report; Backlog and Sprint Plan.',
    'Harborline project repository. Testing and Acceptance Plan, verification results dated 6 October 2026.'
]
for i, p in enumerate(reference_paras[:4]):
    replace(p, '[' + str(i + 1) + '] ' + refs[i], 'Body Text', False)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.runs[-1].font.size = Pt(8.5)

# Restart subsection labels for each chapter while preserving the template's heading appearance.
for label, p in [
    ('a.', method_sub), ('b.', arch_h), ('c.', auth_h),
    ('a.', experiments_sub), ('b.', evaluation_h), ('a.', results_sub)
]:
    clean = p.text.strip()
    replace(p, label + ' ' + clean, 'Body Text', False)
    if p._p.pPr is not None and p._p.pPr.numPr is not None:
        p._p.pPr.remove(p._p.pPr.numPr)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.30)
    p.paragraph_format.first_line_indent = Inches(-0.30)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    p.runs[-1].font.size = Pt(11.5)
    p.runs[-1].italic = True

# Harborline diagrams drawn at the template's existing image placements and proportions.
F_REG = 'C:\\Windows\\Fonts\\times.ttf'
F_BOLD = 'C:\\Windows\\Fonts\\timesbd.ttf'
if not Path(F_REG).exists():
    F_REG = 'C:\\Windows\\Fonts\\arial.ttf'
if not Path(F_BOLD).exists():
    F_BOLD = 'C:\\Windows\\Fonts\\arialbd.ttf'
NAVY, TEAL, PALE = '#12343D', '#0A817D', '#E8F3EF'
BORDER, INK, MUTED = '#BCD0CA', '#1D3338', '#536C71'
WHITE, BG = '#FFFFFF', '#F5F8F6'

def ft(size, bold=False):
    try:
        return ImageFont.truetype(F_BOLD if bold else F_REG, size)
    except Exception:
        return ImageFont.load_default()

def rounded(d, box, fill=WHITE, outline=BORDER, radius=20, width=3):
    d.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)

def wrap(d, text, f, width):
    words, lines, line = text.split(), [], ''
    for word in words:
        test = word if not line else line + ' ' + word
        if d.textlength(test, font=f) <= width or not line:
            line = test
        else:
            lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines

def wrapped(d, x, y, text, f, width, fill=INK, gap=4):
    for line in wrap(d, text, f, width):
        d.text((x, y), line, font=f, fill=fill)
        y += f.size + gap
    return y

def arrow(d, x1, y1, x2, y2, color=TEAL, width=6):
    import math
    d.line((x1, y1, x2, y2), fill=color, width=width)
    a = math.atan2(y2-y1, x2-x1)
    sz = 19
    d.polygon([(x2,y2),(x2-sz*math.cos(a-.45),y2-sz*math.sin(a-.45)),
               (x2-sz*math.cos(a+.45),y2-sz*math.sin(a+.45))], fill=color)

def figure1(path):
    im = Image.new('RGB', (1600, 662), BG); d = ImageDraw.Draw(im)
    d.text((55, 20), 'Harborline application architecture', font=ft(46, True), fill=NAVY)
    d.text((58, 73), 'Current local prototype with optional external payment configuration', font=ft(26), fill=MUTED)
    roles = ['Passenger', 'Crew', 'Captain', 'Manager']
    for i, name in enumerate(roles):
        x = 60 + i*380
        rounded(d, (x, 120, x+350, 184), fill=WHITE, radius=16, width=2)
        bb=d.textbbox((0,0),name,font=ft(27,True))
        d.text((x+175-(bb[2]-bb[0])/2,137),name,font=ft(27,True),fill=NAVY)
    d.text((62, 195), 'Role-scoped browser access', font=ft(21), fill=MUTED)
    boxes = [(55,250,440,'Web client','HTML, CSS and JavaScript'),
             (580,250,440,'Node.js HTTP API','Same-origin JSON routes'),
             (1105,250,440,'Authorization','Role, account and sailing scope')]
    for x,y,w,title,sub in boxes:
        rounded(d,(x,y,x+w,y+100),fill=WHITE,radius=18)
        d.text((x+20,y+12),title,font=ft(29,True),fill=NAVY)
        d.text((x+20,y+57),sub,font=ft(21),fill=MUTED)
    arrow(d,495,300,570,300); arrow(d,1020,300,1095,300)
    modules=[('Bookings','Cruises and cabins'),('Dining','Batches and guest seats'),
             ('Crew work','Tasks and activity logs'),('Ship operations','Ports, muster, maintenance')]
    for i,(title,sub) in enumerate(modules):
        x=55+i*380
        rounded(d,(x,395,x+350,478),fill=PALE,radius=16,width=2)
        d.text((x+14,405),title,font=ft(24,True),fill=NAVY)
        d.text((x+14,442),sub,font=ft(18),fill=MUTED)
    rounded(d,(180,505,745,562),fill=WHITE,radius=14,width=2)
    d.text((198,519),'SQLite  |  persistent local records and audit history',font=ft(20,True),fill=NAVY)
    rounded(d,(790,505,1418,562),fill='#F6F7E7',outline='#D6DDA4',radius=14,width=2)
    d.text((806,519),'Optional Stripe Checkout  |  keys required; demo takes no payment',font=ft(19),fill=INK)
    d.text((55,588),'Figure 1. Harborline architecture and configured integration boundary',font=ft(21),fill=INK)
    im.save(path,optimize=True)

def figure2(path):
    im=Image.new('RGB',(1500,732),BG); d=ImageDraw.Draw(im)
    d.text((45,20),'Role workspaces and key workflows',font=ft(43,True),fill=NAVY)
    cards=[
        ('Passenger','Plan and manage the booking',['Browse sailings and choose a cabin','Assign one seat per ticketed guest','Review itinerary notices and service cases']),
        ('Crew','Complete assigned shipboard work',['View shifts and record attendance','Accept faults and guest-service tasks','Log progress, maintenance and resolution']),
        ('Captain','Direct sailing operations',['Create and dispatch fault reports','Review port changes and recovery plans','Lead muster and operational response']),
        ('Manager','Configure and oversee sailings',['Set dates, fares, cabins and meal batches','Grant sailing-scoped staff permissions','Review activity, feedback and reports'])
    ]
    for (title,desc,items),(x,y) in zip(cards,[(45,95),(770,95),(45,385),(770,385)]):
        rounded(d,(x,y,x+685,y+250),fill=WHITE,radius=22)
        d.rounded_rectangle((x+1,y+1,x+684,y+62),radius=20,fill=PALE)
        d.text((x+22,y+12),title,font=ft(31,True),fill=NAVY)
        d.text((x+22,y+76),desc,font=ft(23,True),fill=TEAL)
        yy=y+116
        for item in items:
            d.ellipse((x+24,yy+6,x+36,yy+18),fill=TEAL)
            wrapped(d,x+49,yy,item,ft(19),610,INK,2)
            yy+=39
    d.text((45,688),'Figure 2. Responsibilities represented in the local passenger and operations prototype',font=ft(20),fill=INK)
    im.save(path,optimize=True)

def figure3(path):
    im=Image.new('RGB',(1600,368),BG); d=ImageDraw.Draw(im)
    d.text((38,15),'Captain-led dispatch and recovery control flow',font=ft(38,True),fill=NAVY)
    steps=[('1  DEFINE','Captain records fault or port disruption'),
           ('2  REVIEW','System scopes sailing, bookings and shore plans'),
           ('3  RESPOND','Crew claims work; passengers record a choice'),
           ('4  LOG','Decisions, progress and closure are retained')]
    for i,(title,body) in enumerate(steps):
        x=45+i*388
        rounded(d,(x,88,x+360,275),fill=PALE if i==1 else WHITE,radius=20)
        d.text((x+18,105),title,font=ft(29,True),fill=TEAL)
        wrapped(d,x+18,158,body,ft(21),325,INK,4)
        if i<3: arrow(d,x+362,180,x+382,180,width=5)
    d.text((45,305),'Figure 3. Human approval connects operational impact to passenger and crew follow-up',font=ft(20),fill=INK)
    im.save(path,optimize=True)

def tables_image(path):
    im=Image.new('RGB',(1500,437),BG); d=ImageDraw.Draw(im)
    cw,gap,x1=700,28,36; x2=x1+cw+gap; top,bottom=18,402
    for x in (x1,x2): rounded(d,(x,top,x+cw,bottom),fill=WHITE,radius=17,width=2)
    def table(x,title,rows):
        d.rounded_rectangle((x+1,top+1,x+cw-1,top+56),radius=16,fill=NAVY)
        d.text((x+17,top+12),title,font=ft(23,True),fill=WHITE)
        y=top+66; col=172
        d.text((x+15,y),'Area',font=ft(18,True),fill=TEAL)
        d.text((x+col,y),'Status',font=ft(18,True),fill=TEAL)
        d.line((x+14,y+29,x+cw-14,y+29),fill=BORDER,width=2); y+=36
        for i,(a,b) in enumerate(rows):
            if i%2==0: d.rectangle((x+9,y-2,x+cw-9,y+52),fill='#F3F7F5')
            d.text((x+15,y+3),a,font=ft(17,True),fill=NAVY)
            wrapped(d,x+col,y+2,b,ft(16),cw-col-25,INK,1)
            d.line((x+14,y+53,x+cw-14,y+53),fill='#E0E9E5',width=1); y+=55
    table(x1,'Table 1  Implemented workflows',[
        ('Passenger','Cabin booking and guest-level batch seats'),
        ('Crew','Accept, update and log assigned jobs'),
        ('Captain','Dispatch faults and approve recovery'),
        ('Manager','Configure inventory and staff access')])
    table(x2,'Table 2  Integration boundaries',[
        ('Payments','Stripe only when server keys are configured'),
        ('Cruise data','Seeded sailings; no live operator feed'),
        ('Boarding','Code 39/manual check-in; no identity adapter'),
        ('Messages','In-app notices; no email or SMS delivery')])
    d.text((38,408),'Sample records and local infrastructure remain; deployment requires operator validation.',font=ft(16),fill=MUTED)
    im.save(path,optimize=True)

for name, fn in [('image2.png',figure1),('image3.png',figure2),('image4.png',figure3),('image5.png',tables_image)]:
    fn(ART/name)

doc.save(str(WORK))
with zipfile.ZipFile(WORK,'r') as zw:
    changed={'word/document.xml':zw.read('word/document.xml')}
for name in ['image2.png','image3.png','image4.png','image5.png']:
    changed['word/media/'+name]=(ART/name).read_bytes()
with zipfile.ZipFile(SRC,'r') as zin, zipfile.ZipFile(DRAFT,'w',zipfile.ZIP_DEFLATED) as zout:
    for info in zin.infolist():
        zout.writestr(info,changed.get(info.filename,zin.read(info.filename)))
out=Document(str(DRAFT))
print('Draft:',DRAFT)
print('Sections:',len(out.sections),'paragraphs:',len(out.paragraphs),'inline images:',len(out.inline_shapes))
print('SHA-256:',hashlib.sha256(DRAFT.read_bytes()).hexdigest())
